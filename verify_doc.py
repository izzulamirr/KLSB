from docx import Document

doc = Document('uploads/cv/KLSB_054_amirun-harraz-bin-budizaman.docx')

print('Document sections in KLSB_054:')
print('=' * 60)

doc_text = '\n'.join(p.text for p in doc.paragraphs)

print('Checking for sections:')
print('✓ EXPERIENCE SUMMARY:' if 'EXPERIENCE SUMMARY' in doc_text else '✗ EXPERIENCE SUMMARY missing')
print('✓ WORKING EXPERIENCE:' if 'WORKING EXPERIENCE' in doc_text else '✗ WORKING EXPERIENCE missing')
print('✓ No REFERENCES' if 'REFERENCES' not in doc_text else '✗ REFERENCES still present (should be removed)')
print('✓ No INVOLVEMENTS' if 'INVOLVEMENTS' not in doc_text else '✗ INVOLVEMENTS still present (should be removed)')

print()
print('Full section order:')
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if para.runs and para.runs[0].bold and len(text) > 5:
        print(f'  • {text}')
