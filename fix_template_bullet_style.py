"""
Update template to set proper bullet style for working experience section
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

template_path = "KLSB_template_true.docx"

print(f"\n{'='*70}")
print(f"UPDATING TEMPLATE BULLET FORMATTING")
print(f"{'='*70}\n")

doc = Document(template_path)

# Find the paragraph with working_experience_formatted tag
for para in doc.paragraphs:
    if '{{working_experience_formatted}}' in para.text:
        print(f"Found working experience placeholder")
        
        # Set paragraph to left align (not justified)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Set paragraph formatting for better bullet rendering
        para.paragraph_format.left_indent = Inches(0)
        para.paragraph_format.first_line_indent = Inches(0)
        
        print(f"  ✓ Set alignment to LEFT")
        print(f"  ✓ Reset indentation")

# Check tables too
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if '{{working_experience_formatted}}' in para.text:
                    print(f"Found working experience placeholder in table")
                    
                    # Set paragraph to left align
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                    # Reset indents
                    para.paragraph_format.left_indent = Inches(0)
                    para.paragraph_format.first_line_indent = Inches(0)
                    
                    print(f"  ✓ Set alignment to LEFT")
                    print(f"  ✓ Reset indentation")

doc.save(template_path)
print(f"\n✓ Template updated successfully!")
print(f"{'='*70}\n")
