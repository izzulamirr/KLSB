"""
Fix template by ensuring tags are in the correct sections
"""

from docx import Document
import os


def fix_template_sections():
    """Fix template to ensure tags are in correct sections"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"FIXING TEMPLATE SECTIONS")
    print(f"{'='*70}\n")
    
    doc = Document(template_path)
    
    # Find section headings and fix tags
    in_education_section = False
    in_work_section = False
    
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        
        # Detect section headers
        if 'ACADEMIC' in text.upper() or 'TECHNICAL QUALIFICATIONS' in text.upper() or 'EDUCATION' in text.upper():
            print(f"Para {idx}: Found EDUCATION section header")
            in_education_section = True
            in_work_section = False
        elif 'WORKING EXPERIENCE' in text.upper() or 'WORK EXPERIENCE' in text.upper():
            print(f"Para {idx}: Found WORK EXPERIENCE section header")
            in_work_section = True
            in_education_section = False
        elif 'OTHERS' in text.upper() or 'TRAINING' in text.upper() or 'SKILLS' in text.upper():
            in_education_section = False
            in_work_section = False
        
        # Fix tags in wrong sections
        if in_work_section and '{{education_formatted}}' in para.text:
            print(f"  ✗ WRONG: Para {idx} in WORK section has {{{{education_formatted}}}}")
            print(f"  → Replacing with {{{{working_experience_formatted}}}}")
            for run in para.runs:
                if '{{education_formatted}}' in run.text:
                    run.text = run.text.replace('{{education_formatted}}', '{{working_experience_formatted}}')
        
        if in_education_section and '{{working_experience_formatted}}' in para.text:
            print(f"  ✗ WRONG: Para {idx} in EDUCATION section has {{{{working_experience_formatted}}}}")
            print(f"  → Replacing with {{{{education_formatted}}}}")
            for run in para.runs:
                if '{{working_experience_formatted}}' in run.text:
                    run.text = run.text.replace('{{working_experience_formatted}}', '{{education_formatted}}')
    
    # Check tables
    for table_idx, table in enumerate(doc.tables):
        # Check if this table is near education section
        # Table 1 should be education
        if table_idx == 1:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if '{{working_experience_formatted}}' in para.text:
                            print(f"  ✗ WRONG: Table {table_idx} (EDUCATION) has {{{{working_experience_formatted}}}}")
                            print(f"  → Replacing with {{{{education_formatted}}}}")
                            for run in para.runs:
                                if '{{working_experience_formatted}}' in run.text:
                                    run.text = run.text.replace('{{working_experience_formatted}}', '{{education_formatted}}')
    
    # Save
    doc.save(template_path)
    print(f"\n✓ Template fixed and saved!")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    fix_template_sections()
