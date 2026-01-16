"""
Inspect template to find klsb_number variable usage
"""

from docx import Document
import os


def inspect_template():
    """Check what variable is used in template for KLSB number"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"INSPECTING TEMPLATE FOR KLSB VARIABLES")
    print(f"{'='*70}\n")
    
    doc = Document(template_path)
    
    found_vars = []
    
    # Check headers
    for section in doc.sections:
        header = section.header
        for para in header.paragraphs:
            text = para.text
            if 'klsb' in text.lower() or '{{' in text:
                print(f"Header paragraph: {text}")
                found_vars.append(text)
        
        # Check header tables
        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        text = para.text
                        if 'klsb' in text.lower() or 'name' in text.lower():
                            print(f"Header table cell: {text}")
                            found_vars.append(text)
    
    # Check main document
    for para in doc.paragraphs:
        text = para.text
        if 'klsb' in text.lower():
            print(f"Body paragraph: {text}")
            found_vars.append(text)
    
    print(f"\n{'='*70}")
    print(f"FOUND {len(found_vars)} REFERENCES")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    inspect_template()
