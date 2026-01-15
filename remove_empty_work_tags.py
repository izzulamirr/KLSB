"""
Remove the individual Year/Company/Position template tags from WORKING EXPERIENCE section
since we're now rendering everything in working_experience_formatted
"""

from docx import Document
import re

def remove_empty_work_experience_tags():
    template_path = "KLSB_template_true.docx"
    
    # Load the document
    doc = Document(template_path)
    
    # Find the WORKING EXPERIENCE section
    work_exp_idx = None
    for idx, para in enumerate(doc.paragraphs):
        if "WORKING EXPERIENCE" in para.text.upper() and para.runs and para.runs[0].bold:
            work_exp_idx = idx
            print(f"✓ Found WORKING EXPERIENCE section at paragraph {idx}")
            break
    
    if work_exp_idx is None:
        print("❌ Could not find WORKING EXPERIENCE section")
        return
    
    # Remove the next 3 paragraphs (Year, Company, Position) after the header
    paragraphs_to_remove = []
    for i in range(1, 4):  # Check next 3 paragraphs
        para_idx = work_exp_idx + i
        if para_idx < len(doc.paragraphs):
            para = doc.paragraphs[para_idx]
            # Check if it contains work_years, work_company, or work_position tags
            if any(tag in para.text for tag in ['{{work_years}}', '{{work_company}}', '{{work_position}}']):
                paragraphs_to_remove.append(para)
                print(f"  → Marking paragraph {para_idx} for removal: {para.text[:50]}...")
    
    # Remove the paragraphs (in reverse order to maintain indices)
    for para in reversed(paragraphs_to_remove):
        # Get the parent element and remove the paragraph
        p_element = para._element
        p_element.getparent().remove(p_element)
        print(f"  ✓ Removed paragraph with text: {para.text[:50]}...")
    
    # Save the updated template
    doc.save(template_path)
    print(f"\n✓ Template updated successfully!")
    print(f"✓ Removed {len(paragraphs_to_remove)} empty tag paragraphs")
    print(f"✓ Changes saved to: {template_path}")

if __name__ == "__main__":
    remove_empty_work_experience_tags()
