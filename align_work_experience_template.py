"""
Update KLSB_template_true.docx to align Year/Company/Position fields using tabs
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
import os


def align_work_experience_fields():
    """Add tab stops to align Year/Company/Position fields"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"ALIGNING WORK EXPERIENCE FIELDS")
    print(f"{'='*70}")
    print(f"Template: {template_path}\n")
    
    try:
        doc = Document(template_path)
        
        changes_made = 0
        
        # Find WORKING EXPERIENCE section
        for para_idx, para in enumerate(doc.paragraphs):
            if 'WORKING EXPERIENCE' in para.text.upper():
                print(f"✓ Found WORKING EXPERIENCE section at paragraph {para_idx}")
                
                # Check next few paragraphs for Year/Company/Position
                for i in range(para_idx + 1, min(para_idx + 10, len(doc.paragraphs))):
                    check_para = doc.paragraphs[i]
                    text = check_para.text.strip()
                    
                    # Update Year, Company, Position paragraphs with tab alignment
                    if any(text.startswith(label) for label in ['Year', 'Company', 'Position']) and ':' in text:
                        print(f"  → Aligning paragraph {i}: {text[:40]}")
                        
                        # Clear existing content
                        for run in check_para.runs:
                            run.text = ''
                        
                        # Set tab stop at 1.5 inches for alignment
                        tab_stops = check_para.paragraph_format.tab_stops
                        tab_stops.add_tab_stop(Inches(1.5))
                        
                        # Determine label
                        if text.startswith('Year'):
                            label = 'Year'
                            tag = '{{work_years}}'
                        elif text.startswith('Company'):
                            label = 'Company'
                            tag = '{{work_company}}'
                        elif text.startswith('Position'):
                            label = 'Position'
                            tag = '{{work_position}}'
                        
                        # Add formatted text with tab
                        run = check_para.add_run(label)
                        run = check_para.add_run('\t:\t')
                        run = check_para.add_run(tag)
                        
                        changes_made += 1
                
                break
        
        # Save
        doc.save(template_path)
        
        print(f"\n✓ Template aligned successfully!")
        print(f"  Changes made: {changes_made}")
        print(f"\n{'='*70}\n")
        
    except PermissionError:
        print(f"\n❌ Permission denied!")
        print(f"   Please close {template_path} in Word and try again.\n")
    except Exception as e:
        print(f"\n❌ Error: {str(e)}\n")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    align_work_experience_fields()
