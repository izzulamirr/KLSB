"""
Check WORKING EXPERIENCE section for any remaining template tags
"""

from docx import Document

def check_work_section():
    template_path = "KLSB_template_true.docx"
    
    # Load the document
    doc = Document(template_path)
    
    # Find the WORKING EXPERIENCE section
    work_exp_idx = None
    for idx, para in enumerate(doc.paragraphs):
        if "WORKING EXPERIENCE" in para.text.upper() and para.runs and para.runs[0].bold:
            work_exp_idx = idx
            print(f"✓ Found WORKING EXPERIENCE section at paragraph {idx}")
            break
    
    if work_exp_idx is None:
        print("❌ Could not find WORKING EXPERIENCE section")
        return
    
    # Check the next 10 paragraphs
    print("\nNext 10 paragraphs after WORKING EXPERIENCE:")
    for i in range(1, 11):
        para_idx = work_exp_idx + i
        if para_idx < len(doc.paragraphs):
            para = doc.paragraphs[para_idx]
            print(f"  Paragraph {para_idx}: {repr(para.text[:100])}")

if __name__ == "__main__":
    check_work_section()
