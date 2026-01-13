"""
Check what tags are in the KLSB_template_true.docx template
"""

from docx import Document
import os


def check_template_tags():
    """Check all tags in the template"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"CHECKING TEMPLATE TAGS")
    print(f"{'='*70}\n")
    
    doc = Document(template_path)
    
    tags_found = []
    
    # Check paragraphs
    for idx, para in enumerate(doc.paragraphs):
        if '{{' in para.text and '}}' in para.text:
            print(f"Para {idx}: {para.text[:100]}")
            tags_found.append(('paragraph', idx, para.text))
    
    # Check tables
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                for para_idx, para in enumerate(cell.paragraphs):
                    if '{{' in para.text and '}}' in para.text:
                        print(f"Table {table_idx}, Row {row_idx}, Cell {cell_idx}: {para.text[:100]}")
                        tags_found.append(('table', (table_idx, row_idx, cell_idx), para.text))
    
    print(f"\n{'='*70}\n")


if __name__ == "__main__":
    check_template_tags()
