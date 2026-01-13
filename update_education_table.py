"""
Update KLSB_template_true.docx education section to use table iteration
This will create a table with Period/Duration | Description/Details columns
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os


def update_education_table():
    """Update education section to use table with period and description"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"UPDATING EDUCATION TABLE IN TEMPLATE")
    print(f"{'='*70}\n")
    
    doc = Document(template_path)
    
    # Find the education table (Table 1 based on previous check)
    if len(doc.tables) < 2:
        print(f"❌ Not enough tables in template")
        return
    
    edu_table = doc.tables[1]
    
    print(f"Found education table with {len(edu_table.rows)} rows")
    
    # Clear existing content except header
    # Keep first row as header if it exists
    header_row = None
    if len(edu_table.rows) > 0:
        first_row = edu_table.rows[0]
        # Check if it's a header (contains "Period" or "Duration")
        first_row_text = ' '.join(cell.text for cell in first_row.cells)
        if 'Period' in first_row_text or 'Duration' in first_row_text or 'Description' in first_row_text:
            print(f"  → Keeping header row")
            header_row = first_row
            # Remove other rows
            for _ in range(len(edu_table.rows) - 1):
                edu_table._element.remove(edu_table.rows[-1]._element)
        else:
            # No header, clear all and add header
            for _ in range(len(edu_table.rows)):
                edu_table._element.remove(edu_table.rows[-1]._element)
    
    # Add header if not exists
    if header_row is None:
        print(f"  → Adding header row")
        header_row = edu_table.add_row()
        header_row.cells[0].text = "Period/Duration"
        header_row.cells[1].text = "Description/Details"
        
        # Make header italic
        for cell in header_row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.italic = True
    
    # Add template tag row for iteration
    print(f"  → Adding template tag row for iteration")
    tag_row = edu_table.add_row()
    
    # Period cell
    tag_row.cells[0].text = ""
    para = tag_row.cells[0].paragraphs[0]
    para.text = "{% for edu in education_table %}{{edu.period}}{% endfor %}"
    
    # Description cell  
    tag_row.cells[1].text = ""
    para = tag_row.cells[1].paragraphs[0]
    para.text = "{% for edu in education_table %}{{edu.description}}{% endfor %}"
    
    print(f"  ✓ Updated education table with iteration tags")
    
    # Save
    doc.save(template_path)
    
    print(f"\n✓ Template updated!")
    print(f"\nEducation table now uses:")
    print(f"  • {{{{education_table}}}} - List of education entries")
    print(f"  • Each entry has: period, description")
    print(f"\n{'='*70}\n")


if __name__ == "__main__":
    update_education_table()
