from docx import Document

doc = Document('KLSB_template.docx')
header = doc.sections[0].header

print('Current Template Header Format:')
print('=' * 60)

# Check if it has tables
if header.tables:
    print(f'Tables in header: {len(header.tables)}')
    for t_idx, table in enumerate(header.tables):
        print(f'\nTable {t_idx + 1}: {len(table.rows)} rows x {len(table.columns)} cols')
        for r_idx, row in enumerate(table.rows):
            print(f'  Row {r_idx + 1}:')
            for c_idx, cell in enumerate(row.cells):
                text = cell.text.strip()
                if text:
                    print(f'    Cell [{c_idx}]: {text}')

# Check paragraphs
print(f'\nParagraphs in header: {len(header.paragraphs)}')
for i, para in enumerate(header.paragraphs):
    text = para.text.strip()
    if text and len(text) < 200:
        align_map = {0: 'LEFT', 1: 'CENTER', 2: 'RIGHT'}
        align = align_map.get(para.alignment, 'NONE')
        print(f'Para {i}: [{align}] {text[:80]}')

print('=' * 60)
