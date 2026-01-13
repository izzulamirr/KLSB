"""
Update education table to use simple text tags (no iteration)
"""

from docx import Document
import os


def update_education_simple_tags():
    """Update education table to use simple text tags"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"UPDATING EDUCATION TABLE - SIMPLE TAGS")
    print(f"{'='*70}\n")
    
    doc = Document(template_path)
    
    # Find education table (Table 1)
    if len(doc.tables) < 2:
        print(f"❌ Not enough tables in template")
        return
    
    edu_table = doc.tables[1]
    
    print(f"Found education table with {len(edu_table.rows)} rows")
    
    # Find the row with template tags
    tag_row_idx = None
    for idx, row in enumerate(edu_table.rows):
        row_text = ' '.join(cell.text for cell in row.cells)
        if '{%' in row_text or '{{' in row_text:
            tag_row_idx = idx
            print(f"  → Found tag row at index {idx}")
            break
    
    if tag_row_idx is None:
        print(f"❌ Could not find template tag row")
        return
    
    # Update the tag row with simple tags
    tag_row = edu_table.rows[tag_row_idx]
    
    # Clear and update Period cell
    period_cell = tag_row.cells[0]
    period_cell.text = ""
    para = period_cell.paragraphs[0]
    para.text = "{{education_periods}}"
    
    # Clear and update Description cell
    desc_cell = tag_row.cells[1]
    desc_cell.text = ""
    para = desc_cell.paragraphs[0]
    para.text = "{{education_descriptions}}"
    
    print(f"  ✓ Updated with simple tags")
    print(f"    Period cell: {{{{education_periods}}}}")
    print(f"    Description cell: {{{{education_descriptions}}}}")
    
    # Save
    doc.save(template_path)
    
    print(f"\n✓ Template updated!")
    print(f"\nEach education entry will be separated by blank lines in table cells")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    update_education_simple_tags()
