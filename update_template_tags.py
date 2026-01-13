"""
Update KLSB_template_true.docx to use the new formatted tags:
- {{education_formatted}} instead of individual education tags
- {{working_experience_formatted}} instead of individual work experience tags
"""

from docx import Document
import os


def update_template_tags():
    """Update template tags in the document"""
    
    template_path = "KLSB_template_true.docx"
    
    if not os.path.exists(template_path):
        print(f"❌ Template not found: {template_path}")
        return
    
    print(f"\n{'='*70}")
    print(f"UPDATING TEMPLATE TAGS")
    print(f"{'='*70}")
    print(f"Template: {template_path}\n")
    
    # Load document
    doc = Document(template_path)
    
    replacements = {
        '{{edu_years}}': '{{education_formatted}}',
        '{{edu_university}}': '{{education_formatted}}',
        '{{edu_degree}}': '{{education_formatted}}',
        '{{edu_description}}': '{{education_formatted}}',
        '{{education}}': '{{education_formatted}}',
        '{{work_years}}': '{{working_experience_formatted}}',
        '{{work_company}}': '{{working_experience_formatted}}',
        '{{work_position}}': '{{working_experience_formatted}}',
        '{{work_description}}': '{{working_experience_formatted}}',
        '{{working_experience}}': '{{working_experience_formatted}}',
    }
    
    replacements_made = {}
    
    # Update paragraphs
    for para in doc.paragraphs:
        for old_tag, new_tag in replacements.items():
            if old_tag in para.text:
                print(f"  → Found in paragraph: {old_tag}")
                # Replace in the paragraph
                if old_tag not in replacements_made:
                    replacements_made[old_tag] = 0
                replacements_made[old_tag] += 1
                
                # Update runs (text segments)
                for run in para.runs:
                    if old_tag in run.text:
                        run.text = run.text.replace(old_tag, new_tag)
    
    # Update tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old_tag, new_tag in replacements.items():
                        if old_tag in para.text:
                            print(f"  → Found in table: {old_tag}")
                            if old_tag not in replacements_made:
                                replacements_made[old_tag] = 0
                            replacements_made[old_tag] += 1
                            
                            for run in para.runs:
                                if old_tag in run.text:
                                    run.text = run.text.replace(old_tag, new_tag)
    
    # Save updated template
    doc.save(template_path)
    
    print(f"\n✓ Template updated successfully!")
    print(f"\nReplacements made:")
    for old_tag, count in replacements_made.items():
        print(f"  {old_tag}: {count} occurrence(s)")
    
    if not replacements_made:
        print("  No tags found to replace (tags might already be updated)")
    
    print(f"\n✓ New template tags available:")
    print(f"  • {{{{education_formatted}}}} - All education entries")
    print(f"  • {{{{working_experience_formatted}}}} - All work experience entries")
    print(f"\n{'='*70}\n")


if __name__ == "__main__":
    update_template_tags()
